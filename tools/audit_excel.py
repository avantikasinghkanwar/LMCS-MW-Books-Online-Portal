#!/usr/bin/env python3
"""Checks the portal against the M1 course-map Excel (the master copy) and lists every difference.

Run from the portal folder:  python3 tools/audit_excel.py
Exit code 0 = portal matches the Excel; 1 = differences found.

What it compares
  letters   : title, the 3 words, Hindi cue, Montessori activity, materials, listening-speaking task, sound tip
  pages     : for both books, the order of pages and their type (letter / pre-writing / practice / chart / revision),
              the letter or skill, the new word on each picture page, and the printed instruction when the Excel gives one
What it does not compare: emoji pictures (placeholders) and audio.
"""
import json
import os
import re
import subprocess
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.expanduser('~/Desktop/LMCS MW Books/01_Course_Mapping/English/M1/English_M1_Course_Mapping.xlsx')


def portal():
    js = ('ObjC.import("Foundation");'
          'function rd(f){return $.NSString.stringWithContentsOfFileEncodingError(f,4,null).js}'
          'eval(rd("content.js")+";this.C=CONTENT;"+rd("activities.js")+";this.A=ACTIVITIES"); JSON.stringify({C:C,A:A})')
    out = subprocess.run(['osascript', '-l', 'JavaScript', '-e', js], cwd=ROOT, capture_output=True, text=True, check=True).stdout
    return json.loads(out)


def clean(v):
    return ' '.join(str(v).split()) if v else ''


def excel_letters(wb):
    ws = wb['M1 English']
    rows = list(ws.iter_rows())
    hdr = [c.value for c in rows[0]]
    col = {h: i for i, h in enumerate(hdr)}
    out = {}
    for r in rows[1:]:
        k = r[col['Letter (Capital)']].value
        if not k or len(str(k)) != 1:
            continue
        cue_cell = r[col['Hindi Cue — Letter Sound (teacher bridge)']]
        out[k] = {
            'title': clean(r[col['Chapter / Topic Title']].value),
            'words': [w.strip().upper() for w in str(r[col['Vocabulary (English)']].value).split(',')],
            'cue': clean(cue_cell.value),
            'montessori': clean(r[col['Montessori Activity — Vocabulary/Practical Life']].value),
            'materials': clean(r[col['Materials Needed']].value),
            'speaking': clean(r[col['Listening-Speaking Activity']].value),
            'sound': clean(cue_cell.comment.text.split('\n', 1)[-1]) if cue_cell.comment else '',
        }
    return out


def excel_items(wb, sheet):
    """Turn the page-by-page rows into the same kind of items the portal uses."""
    rows = [[c.value for c in r] for r in wb[sheet].iter_rows()][1:]
    rows = [r for r in rows if isinstance(r[0], int)]
    items, i = [], 0
    while i < len(rows):
        r = rows[i]
        t = r[2]
        nxt = rows[i + 1] if i + 1 < len(rows) else None
        if t == 'A–Z Recognition Chart':
            items.append({'type': 'chart', 'pages': [r[0]]}); i += 1
        elif t == 'Letter Intro/Tracing':
            items.append({'type': 'letter', 'letter': r[3], 'pages': [r[0], nxt[0]]}); i += 2
        elif t == 'Alphabet Colouring (Revision)':
            items.append({'type': 'colourletters', 'pages': [r[0], nxt[0]]}); i += 2
        elif t == 'Blank Colouring':
            items.append({'type': 'blank', 'pages': [r[0], nxt[0]]}); i += 2
        elif t == 'Pre-Writing Skill':
            if r[10] == 'Finger Tracing':      # picture page + its practice page
                m = re.search(r'Instruction: "([^"]+)"', r[9] or '')
                items.append({'type': 'prewriting', 'skill': r[5], 'word': (r[7] or '').strip(), 'say': m.group(1) if m else None,
                              'pages': [r[0], nxt[0]]}); i += 2
            else:                              # two practice-only pages, each with its own shapes
                items.append({'type': 'tracepractice', 'skill': r[5] + ' | ' + nxt[5], 'pages': [r[0], nxt[0]]}); i += 2
        else:
            items.append({'type': 'UNKNOWN:' + str(t), 'pages': [r[0]]}); i += 1
    return items


def main():
    wb = openpyxl.load_workbook(XLSX)
    p = portal()
    C, A = p['C'], p['A']
    problems = []

    # ---- letters ----
    ex = excel_letters(wb)
    for k, e in ex.items():
        L = C['LETTERS'].get(k)
        if not L:
            problems.append('letter %s: missing in portal' % k); continue
        if L['title'] != e['title']:
            problems.append('letter %s: title  Excel=%r  portal=%r' % (k, e['title'], L['title']))
        pw = [w['name'] for w in L['words']]
        if pw != e['words']:
            problems.append('letter %s: words  Excel=%s  portal=%s' % (k, e['words'], pw))
        if not L['cue'].startswith(e['cue']):
            problems.append('letter %s: cue  Excel=%r  portal=%r' % (k, e['cue'], L['cue']))
        a = A.get(k, {})
        for f in ('montessori', 'materials', 'speaking', 'sound'):
            if a.get(f, '') != e[f]:
                problems.append('letter %s: activity field "%s" differs (run tools/import_activities.py)' % (k, f))
    for k in C['LETTERS']:
        if k not in ex:
            problems.append('letter %s: in portal but not in Excel' % k)

    # ---- pages ----
    skills = {k: v for k, v in C['SKILLS'].items()}
    by_excel = {v.get('excel'): k for k, v in skills.items()}
    for bk, sheet in (('1', 'Book 1 Page Plan (A-N)'), ('2', 'Book 2 Page Plan (O-Z)')):
        exi = excel_items(wb, sheet)
        pi = C['BOOKS'][bk]['items']
        tag = 'Book %s' % bk
        if len(exi) != len(pi):
            problems.append('%s: Excel has %d page groups, portal has %d' % (tag, len(exi), len(pi)))
        for n, (e, q) in enumerate(zip(exi, pi)):
            where = '%s pages %s' % (tag, '–'.join(str(x) for x in e['pages']))
            if e['pages'] != q['pages'] or e['type'] != q['type']:
                problems.append('%s: Excel=%s %s  portal=%s %s' % (where, e['type'], e['pages'], q['type'], q['pages'])); continue
            if e['type'] == 'letter' and e['letter'] != q['letter']:
                problems.append('%s: letter Excel=%s portal=%s' % (where, e['letter'], q['letter']))
            if e['type'] == 'prewriting':
                if skills.get(q['skill'], {}).get('excel') != e['skill']:
                    problems.append('%s: skill  Excel=%r  portal=%r' % (where, e['skill'], skills.get(q['skill'], {}).get('excel')))
                if (q.get('word') or '') != e['word']:
                    problems.append('%s: new word  Excel=%r  portal=%r' % (where, e['word'], q.get('word') or ''))
                if e['say'] and q.get('say') != e['say']:
                    problems.append('%s: instruction  Excel=%r  portal=%r' % (where, e['say'], q.get('say')))
            if e['type'] == 'tracepractice':
                want = [s.strip() for s in e['skill'].split('|')]
                got = ['%s + %s' % tuple(skills[x]['excel'] for x in pair) for pair in q.get('sets', [])]
                exp = ['%s + %s' % tuple(y.strip() for y in w.split('+')[:1] + w.split('+')[1:]) for w in want] if False else want
                if got != exp:
                    problems.append('%s: practice shapes  Excel=%s  portal=%s' % (where, exp, got))

    if problems:
        print('%d difference(s) between the Excel (master) and the portal:' % len(problems))
        for x in problems:
            print('  -', x)
        sys.exit(1)
    print('OK: the portal matches the Excel — %d letters, Book 1 (%d groups), Book 2 (%d groups).' % (
        len(ex), len(C['BOOKS']['1']['items']), len(C['BOOKS']['2']['items'])))


if __name__ == '__main__':
    main()
